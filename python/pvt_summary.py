import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter


# =============================================================================
# Project Paths
# =============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

RAW_DIR = PROJECT_ROOT / "data" / "raw"
RESULT_DIR = PROJECT_ROOT / "results"

SETUP_FILE = RAW_DIR / "SETUP_TIME_PVT.csv"
HOLD_FILE = RAW_DIR / "HOLD_TIME_PVT.csv"
ETC_FILE = RAW_DIR / "ETC_PVT.csv"

EXCEL_FILE = RESULT_DIR / "PVT_RESULT.xlsx"

RESULT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


# =============================================================================
# Report Theme
# =============================================================================

COLOR_NAVY = "0F172A"
COLOR_SLATE = "1E293B"
COLOR_SECTION = "334155"
COLOR_WHITE = "FFFFFF"
COLOR_ROW_ALT = "F1F5F9"
COLOR_BORDER = "CBD5E1"
COLOR_TEXT = "1E293B"

FONT_NAME = "Aptos"

SUMMARY_MARKER = "========== STATISTICS =========="
SUMMARY_TITLE = "STATISTICS SUMMARY"

STAT_HEADERS = [
    "Metric",
    "MIN",
    "MIN CONDITION",
    "MAX",
    "MAX CONDITION",
    "MEAN",
]


# =============================================================================
# General Utilities
# =============================================================================

def to_float(value):
    """Convert a CSV value to float, returning None for invalid values."""

    if value is None:
        return None

    value = value.strip()

    if value == "" or value.upper() == "N/A":
        return None

    try:
        return float(value)

    except ValueError:
        return None


def format_number(value):
    """Format a statistic with six significant digits."""

    if value is None:
        return "N/A"

    return f"{value:.6g}"


def humanize_metric_name(metric_name):
    """Convert internal metric identifiers into report-friendly names."""

    display_name_map = {
        "SETUP_TIME(ps)": "Setup Time",
        "HOLD_TIME(ps)": "Hold Time",
        "SETUP_TIME": "Setup Time",
        "HOLD_TIME": "Hold Time",
        "DELAY_CLK_TO_Q": "CLK-to-Q Delay",
        "CLK_TO_Q_DELAY": "CLK-to-Q Delay",
        "Q_RISING_TIME": "Output Rise Time",
        "Q_FALLING_TIME": "Output Fall Time",
        "OUTPUT_RISE_TIME": "Output Rise Time",
        "OUTPUT_FALL_TIME": "Output Fall Time",
        "POWER_CONSUMPTION": "Power Consumption",
    }

    return display_name_map.get(
        metric_name,
        metric_name,
    )


def read_raw_rows(filename):
    """
    Read one OCEAN-generated CSV.

    Legacy statistics appended by an older Python version are ignored so the
    script can also process previously generated CSV files safely.
    """

    with open(
        filename,
        "r",
        newline="",
    ) as file_handle:

        rows = list(
            csv.reader(file_handle)
        )

    raw_rows = []

    for row in rows:

        if row and row[0] == SUMMARY_MARKER:
            break

        raw_rows.append(
            row
        )

    while raw_rows and not raw_rows[-1]:
        raw_rows.pop()

    return raw_rows


def calc_stats(data):
    """
    Calculate min, max, and mean.

    data:
        [(value, condition), ...]
    """

    if not data:
        return None

    min_item = min(
        data,
        key=lambda item: item[0],
    )

    max_item = max(
        data,
        key=lambda item: item[0],
    )

    mean_value = (
        sum(
            item[0]
            for item in data
        )
        /
        len(data)
    )

    return {
        "min": min_item[0],
        "min_condition": min_item[1],
        "max": max_item[0],
        "max_condition": max_item[1],
        "mean": mean_value,
    }


def make_stat_row(
    metric_name,
    stats,
):
    """Create one statistics-table row."""

    if stats is None:

        return [
            humanize_metric_name(metric_name),
            "N/A",
            "",
            "N/A",
            "",
            "N/A",
        ]

    return [
        humanize_metric_name(metric_name),
        format_number(
            stats["min"]
        ),
        stats["min_condition"],
        format_number(
            stats["max"]
        ),
        stats["max_condition"],
        format_number(
            stats["mean"]
        ),
    ]


# =============================================================================
# Statistics Extraction
# =============================================================================

def collect_timing_statistics(
    rows,
    metric_name,
):
    """Collect one Setup/Hold metric across all PVT conditions."""

    data = []

    current_process = None
    temperature_headers = []

    for row in rows:

        if not row:
            continue

        first_value = row[0].strip()

        if first_value == "Process":

            current_process = (
                row[1].strip()
                if len(row) > 1
                else ""
            )

            continue

        if first_value == "VDD(V)":

            temperature_headers = [
                item.strip()
                for item in row[1:]
            ]

            continue

        if (
            current_process is None
            or
            not temperature_headers
        ):
            continue

        vdd = first_value

        for temperature, value in zip(
            temperature_headers,
            row[1:],
        ):

            number = to_float(
                value
            )

            if number is None:
                continue

            condition = (
                f"{current_process} / "
                f"{vdd} V / "
                f"{temperature}"
            )

            data.append(
                (
                    number,
                    condition,
                )
            )

    return [
        make_stat_row(
            metric_name,
            calc_stats(data),
        )
    ]


def collect_etc_statistics(rows):
    """Collect every user-defined ETC scalar metric across all PVT conditions."""

    metric_data = {}

    current_process = None
    condition_headers = []

    for row in rows:

        if not row:
            continue

        first_value = row[0].strip()

        if first_value == "Process":

            current_process = (
                row[1].strip()
                if len(row) > 1
                else ""
            )

            continue

        if first_value == "Metric":

            condition_headers = [
                item.strip()
                for item in row[1:]
            ]

            continue

        if (
            not first_value
            or
            current_process is None
            or
            not condition_headers
        ):
            continue

        metric_data.setdefault(
            first_value,
            [],
        )

        for condition_part, value in zip(
            condition_headers,
            row[1:],
        ):

            number = to_float(
                value
            )

            if number is None:
                continue

            condition = (
                f"{current_process} / "
                f"{condition_part}"
            )

            metric_data[
                first_value
            ].append(
                (
                    number,
                    condition,
                )
            )

    statistics_rows = []

    for metric_name, data in metric_data.items():

        statistics_rows.append(
            make_stat_row(
                metric_name,
                calc_stats(data),
            )
        )

    return statistics_rows


# =============================================================================
# Excel Styling Helpers
# =============================================================================

def apply_outer_border(
    ws,
    min_row,
    max_row,
    min_col,
    max_col,
    border_side,
):
    """Apply one continuous outer border around a rectangular range."""

    for row_index in range(
        min_row,
        max_row + 1,
    ):

        for column_index in range(
            min_col,
            max_col + 1,
        ):

            cell = ws.cell(
                row=row_index,
                column=column_index,
            )

            left = cell.border.left
            right = cell.border.right
            top = cell.border.top
            bottom = cell.border.bottom

            if column_index == min_col:
                left = border_side

            if column_index == max_col:
                right = border_side

            if row_index == min_row:
                top = border_side

            if row_index == max_row:
                bottom = border_side

            cell.border = Border(
                left=left,
                right=right,
                top=top,
                bottom=bottom,
            )


def find_process_blocks(
    ws,
    start_row,
    end_row,
    stats_row,
):
    """Find each Process block in the displayed PVT table."""

    process_rows = []

    for row_index in range(
        start_row,
        stats_row,
    ):

        if (
            ws.cell(
                row=row_index,
                column=1,
            ).value
            ==
            "Process"
        ):

            process_rows.append(
                row_index
            )

    blocks = []

    for index, process_row in enumerate(
        process_rows
    ):

        if index + 1 < len(process_rows):

            block_end = (
                process_rows[index + 1]
                -
                1
            )

        else:

            block_end = (
                stats_row
                -
                2
            )

        while (
            block_end > process_row
            and
            ws.cell(
                row=block_end,
                column=1,
            ).value
            in (
                None,
                "",
            )
        ):
            block_end -= 1

        blocks.append(
            (
                process_row,
                block_end,
            )
        )

    return blocks


def style_title(
    ws,
    title,
    max_columns,
):
    """Create the report title row."""

    ws.sheet_view.showGridLines = False

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max_columns,
    )

    title_cell = ws.cell(
        row=1,
        column=1,
    )

    title_cell.value = title

    title_cell.font = Font(
        name=FONT_NAME,
        size=16,
        bold=True,
        color=COLOR_WHITE,
    )

    title_cell.fill = PatternFill(
        fill_type="solid",
        fgColor=COLOR_NAVY,
    )

    title_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    ws.row_dimensions[1].height = 30


def style_sheet(
    ws,
    max_columns,
    stats_title_row,
):
    """Apply the common PVT worksheet style."""

    thin_side = Side(
        style="thin",
        color=COLOR_BORDER,
    )

    medium_side = Side(
        style="medium",
        color=COLOR_SLATE,
    )

    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    # -------------------------------------------------------------------------
    # Base Font / Alignment
    # -------------------------------------------------------------------------

    for row in ws.iter_rows(
        min_row=2,
    ):

        for cell in row:

            cell.font = Font(
                name=FONT_NAME,
                size=10,
                color=COLOR_TEXT,
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    # -------------------------------------------------------------------------
    # Process Rows and PVT Headers
    # -------------------------------------------------------------------------

    for row_index in range(
        2,
        stats_title_row,
    ):

        first_value = ws.cell(
            row=row_index,
            column=1,
        ).value

        if first_value == "Process":

            for column_index in range(
                1,
                max_columns + 1,
            ):

                cell = ws.cell(
                    row=row_index,
                    column=column_index,
                )

                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=COLOR_SECTION,
                )

                cell.font = Font(
                    name=FONT_NAME,
                    size=10,
                    bold=True,
                    color=COLOR_WHITE,
                )

                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="center",
                )

            ws.row_dimensions[
                row_index
            ].height = 22

        elif first_value in (
            "VDD(V)",
            "Metric",
        ):

            for column_index in range(
                1,
                max_columns + 1,
            ):

                cell = ws.cell(
                    row=row_index,
                    column=column_index,
                )

                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=COLOR_SLATE,
                )

                cell.font = Font(
                    name=FONT_NAME,
                    size=10,
                    bold=True,
                    color=COLOR_WHITE,
                )

                cell.border = thin_border

            ws.row_dimensions[
                row_index
            ].height = 22

    # -------------------------------------------------------------------------
    # Alternating PVT Data Rows
    # -------------------------------------------------------------------------

    process_blocks = find_process_blocks(
        ws,
        start_row=2,
        end_row=ws.max_row,
        stats_row=stats_title_row,
    )

    for block_start, block_end in process_blocks:

        data_row_counter = 0

        for row_index in range(
            block_start + 1,
            block_end + 1,
        ):

            first_value = ws.cell(
                row=row_index,
                column=1,
            ).value

            if first_value in (
                "VDD(V)",
                "Metric",
                None,
                "",
            ):
                continue

            data_row_counter += 1

            fill_color = (
                COLOR_ROW_ALT
                if data_row_counter % 2 == 0
                else COLOR_WHITE
            )

            fill = PatternFill(
                fill_type="solid",
                fgColor=fill_color,
            )

            for column_index in range(
                1,
                max_columns + 1,
            ):

                cell = ws.cell(
                    row=row_index,
                    column=column_index,
                )

                cell.fill = fill

                cell.border = Border(
                    bottom=thin_side,
                )

    # -------------------------------------------------------------------------
    # Process Block Borders
    # -------------------------------------------------------------------------

    for block_start, block_end in process_blocks:

        apply_outer_border(
            ws,
            min_row=block_start,
            max_row=block_end,
            min_col=1,
            max_col=max_columns,
            border_side=medium_side,
        )

    # -------------------------------------------------------------------------
    # Statistics Table
    # -------------------------------------------------------------------------

    stats_header_row = (
        stats_title_row
        +
        1
    )

    stats_last_row = ws.max_row

    ws.merge_cells(
        start_row=stats_title_row,
        start_column=1,
        end_row=stats_title_row,
        end_column=6,
    )

    stats_title_cell = ws.cell(
        row=stats_title_row,
        column=1,
    )

    stats_title_cell.value = SUMMARY_TITLE

    stats_title_cell.fill = PatternFill(
        fill_type="solid",
        fgColor=COLOR_NAVY,
    )

    stats_title_cell.font = Font(
        name=FONT_NAME,
        size=11,
        bold=True,
        color=COLOR_WHITE,
    )

    stats_title_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    ws.row_dimensions[
        stats_title_row
    ].height = 22

    for column_index in range(
        1,
        7,
    ):

        cell = ws.cell(
            row=stats_header_row,
            column=column_index,
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=COLOR_SLATE,
        )

        cell.font = Font(
            name=FONT_NAME,
            size=10,
            bold=True,
            color=COLOR_WHITE,
        )

        cell.border = thin_border

    for row_index in range(
        stats_header_row + 1,
        stats_last_row + 1,
    ):

        fill_color = (
            COLOR_ROW_ALT
            if (
                row_index
                -
                stats_header_row
            ) % 2 == 0
            else COLOR_WHITE
        )

        fill = PatternFill(
            fill_type="solid",
            fgColor=fill_color,
        )

        for column_index in range(
            1,
            7,
        ):

            cell = ws.cell(
                row=row_index,
                column=column_index,
            )

            cell.fill = fill

            cell.border = Border(
                bottom=thin_side,
            )

            if column_index == 1:

                cell.font = Font(
                    name=FONT_NAME,
                    size=10,
                    bold=True,
                    color=COLOR_NAVY,
                )

    apply_outer_border(
        ws,
        min_row=stats_title_row,
        max_row=stats_last_row,
        min_col=1,
        max_col=6,
        border_side=medium_side,
    )

    # -------------------------------------------------------------------------
    # Column Widths
    # -------------------------------------------------------------------------

    fixed_widths = {
        "A": 28,
        "B": 18,
        "C": 35,
        "D": 18,
        "E": 35,
        "F": 18,
    }

    for column_letter, width in fixed_widths.items():

        ws.column_dimensions[
            column_letter
        ].width = width

    for column_index in range(
        7,
        max_columns + 1,
    ):

        ws.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = 18

    ws.freeze_panes = "A2"


# =============================================================================
# Worksheet Builder
# =============================================================================

def build_sheet(
    ws,
    raw_rows,
    statistics_rows,
    title,
):
    """Write one complete PVT worksheet."""

    max_columns = max(
        6,
        max(
            (
                len(row)
                for row in raw_rows
            ),
            default=6,
        ),
    )

    style_title(
        ws,
        title,
        max_columns,
    )

    # -------------------------------------------------------------------------
    # Raw PVT Data
    # -------------------------------------------------------------------------

    excel_row = 2

    for row in raw_rows:

        for column_index, value in enumerate(
            row,
            start=1,
        ):

            display_value = value

            if column_index == 1:

                display_value = humanize_metric_name(
                    value
                )

            ws.cell(
                row=excel_row,
                column=column_index,
                value=display_value,
            )

        excel_row += 1

    # One blank row between raw characterization data and statistics.
    excel_row += 1

    stats_title_row = excel_row

    # -------------------------------------------------------------------------
    # Statistics Summary
    # -------------------------------------------------------------------------

    for column_index, header in enumerate(
        STAT_HEADERS,
        start=1,
    ):

        ws.cell(
            row=stats_title_row + 1,
            column=column_index,
            value=header,
        )

    stats_data_start = (
        stats_title_row
        +
        2
    )

    for row_offset, row in enumerate(
        statistics_rows
    ):

        for column_index, value in enumerate(
            row,
            start=1,
        ):

            ws.cell(
                row=stats_data_start + row_offset,
                column=column_index,
                value=value,
            )

    style_sheet(
        ws,
        max_columns,
        stats_title_row,
    )


# =============================================================================
# Main
# =============================================================================

def main():

    required_files = [
        SETUP_FILE,
        HOLD_FILE,
        ETC_FILE,
    ]

    for file_path in required_files:

        if not file_path.exists():

            raise FileNotFoundError(
                f"CSV file not found: {file_path}"
            )

    # -------------------------------------------------------------------------
    # Read Raw Characterization Results
    # -------------------------------------------------------------------------

    setup_rows = read_raw_rows(
        SETUP_FILE
    )

    hold_rows = read_raw_rows(
        HOLD_FILE
    )

    etc_rows = read_raw_rows(
        ETC_FILE
    )

    # -------------------------------------------------------------------------
    # Calculate Statistics In Memory
    # -------------------------------------------------------------------------

    setup_statistics = collect_timing_statistics(
        setup_rows,
        "SETUP_TIME(ps)",
    )

    hold_statistics = collect_timing_statistics(
        hold_rows,
        "HOLD_TIME(ps)",
    )

    etc_statistics = collect_etc_statistics(
        etc_rows
    )

    # -------------------------------------------------------------------------
    # Create Excel Workbook
    # -------------------------------------------------------------------------

    workbook = Workbook()

    workbook.remove(
        workbook.active
    )

    sheet_specs = [
        (
            "SETUP",
            setup_rows,
            setup_statistics,
            "PVT SETUP TIME CHARACTERIZATION",
        ),
        (
            "HOLD",
            hold_rows,
            hold_statistics,
            "PVT HOLD TIME CHARACTERIZATION",
        ),
        (
            "ETC",
            etc_rows,
            etc_statistics,
            "PVT ETC CHARACTERIZATION",
        ),
    ]

    for (
        sheet_name,
        raw_rows,
        statistics_rows,
        title,
    ) in sheet_specs:

        worksheet = workbook.create_sheet(
            sheet_name
        )

        build_sheet(
            worksheet,
            raw_rows,
            statistics_rows,
            title,
        )

    workbook.save(
        EXCEL_FILE
    )

    # -------------------------------------------------------------------------
    # Console Summary
    # -------------------------------------------------------------------------

    print("")
    print("===================================================")
    print(" PVT Post Processing Completed")
    print("===================================================")
    print("")
    print("Input CSV:")
    print("  ", SETUP_FILE)
    print("  ", HOLD_FILE)
    print("  ", ETC_FILE)
    print("")
    print("Excel created:")
    print("  ", EXCEL_FILE)
    print("")


if __name__ == "__main__":
    main()

