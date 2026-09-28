import csv
import statistics
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)


# =============================================================================
# Project Paths
# =============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

RAW_DIR = PROJECT_ROOT / "data" / "raw"
RESULT_DIR = PROJECT_ROOT / "results"

SETUP_CSV = RAW_DIR / "SETUP_TIME_MONTE.csv"
HOLD_CSV = RAW_DIR / "HOLD_TIME_MONTE.csv"
ETC_CSV = RAW_DIR / "ETC_MONTE.csv"

OUTPUT_XLSX = RESULT_DIR / "MONTE_RESULT.xlsx"

RESULT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


# =============================================================================
# Report Configuration
# =============================================================================

COLOR_NAVY = "0F172A"
COLOR_SLATE = "1E293B"
COLOR_WHITE = "FFFFFF"
COLOR_ROW_ALT = "F1F5F9"
COLOR_BORDER = "CBD5E1"
COLOR_TEXT = "1E293B"

FONT_NAME = "Aptos"

TABLE_HEADERS = [
    "Metric",
    "Unit",
    "N",
    "Mean",
    "1-Sigma",
    "Min",
    "Max",
]

FRAMEWORK_COLUMNS = {
    "Run",
    "Temperature",
    "VDD",
}

FIXED_METRICS = [
    {
        "display_name": "SETUP_TIME",
        "csv_path": SETUP_CSV,
        "csv_column": "Setup_Time_ps",
        "unit": "ps",
    },
    {
        "display_name": "HOLD_TIME",
        "csv_path": HOLD_CSV,
        "csv_column": "Hold_Time_ps",
        "unit": "ps",
    },
]


# =============================================================================
# Data Utilities
# =============================================================================

def safe_float(value):
    """Convert one CSV field to float. Return None for invalid data."""

    if value is None:
        return None

    value = value.strip()

    if value == "" or value.upper() == "N/A":
        return None

    try:
        return float(value)

    except ValueError:
        return None


def humanize_metric_name(metric_name):
    """Convert an internal metric identifier into a report-friendly name."""

    display_name_map = {
        "SETUP_TIME": "Setup Time",
        "HOLD_TIME": "Hold Time",
        "DELAY_CLK_TO_Q": "CLK-to-Q Delay",
        "CLK_TO_Q_DELAY": "CLK-to-Q Delay",
        "Q_RISING_TIME": "Output Rise Time",
        "Q_FALLING_TIME": "Output Fall Time",
        "OUTPUT_RISE_TIME": "Output Rise Time",
        "OUTPUT_FALL_TIME": "Output Fall Time",
        "POWER_CONSUMPTION": "Power Consumption",
    }

    return display_name_map.get(
        metric_name,
        metric_name.replace("_", " ").title(),
    )


def read_numeric_columns(
    csv_path,
    metric_names=None,
):
    """
    Read selected numeric columns from one Monte Carlo CSV.

    If metric_names is None, every column except the framework metadata
    columns (Run, Temperature, VDD) is treated as a scalar metric.
    """

    with open(
        csv_path,
        "r",
        newline="",
    ) as file_handle:

        reader = csv.DictReader(
            file_handle
        )

        fieldnames = (
            reader.fieldnames
            or
            []
        )

        if metric_names is None:

            metric_names = [
                name
                for name in fieldnames
                if name not in FRAMEWORK_COLUMNS
            ]

        values_by_metric = {
            metric_name: []
            for metric_name in metric_names
        }

        for row in reader:

            for metric_name in metric_names:

                value = safe_float(
                    row.get(
                        metric_name
                    )
                )

                if value is not None:

                    values_by_metric[
                        metric_name
                    ].append(
                        value
                    )

    return values_by_metric


# =============================================================================
# Statistics
# =============================================================================

def calc_statistics(values):
    """
    Calculate pooled Monte Carlo statistics.

    1-Sigma uses the sample standard deviation:
        statistics.stdev()
        denominator = N - 1
        equivalent to Excel STDEV.S
    """

    n = len(
        values
    )

    if n == 0:

        return {
            "N": 0,
            "Mean": None,
            "Sigma": None,
            "Min": None,
            "Max": None,
        }

    sigma_value = (
        statistics.stdev(
            values
        )
        if n >= 2
        else None
    )

    return {
        "N": n,
        "Mean": statistics.mean(
            values
        ),
        "Sigma": sigma_value,
        "Min": min(
            values
        ),
        "Max": max(
            values
        ),
    }


def make_summary_row(
    metric_name,
    unit,
    values,
):
    """Build one Excel summary row."""

    stats = calc_statistics(
        values
    )

    return [
        humanize_metric_name(
            metric_name
        ),
        unit,
        stats["N"],
        stats["Mean"],
        stats["Sigma"],
        stats["Min"],
        stats["Max"],
    ]


def collect_summary_rows():
    """Read all Monte CSV files once and build the report rows."""

    summary_rows = []

    # -------------------------------------------------------------------------
    # Setup / Hold Timing Metrics
    # -------------------------------------------------------------------------

    for metric in FIXED_METRICS:

        values_by_metric = read_numeric_columns(
            metric["csv_path"],
            [
                metric["csv_column"]
            ],
        )

        values = values_by_metric[
            metric["csv_column"]
        ]

        summary_rows.append(
            make_summary_row(
                metric["display_name"],
                metric["unit"],
                values,
            )
        )

    # -------------------------------------------------------------------------
    # User-Defined ETC Metrics
    # -------------------------------------------------------------------------

    etc_values = read_numeric_columns(
        ETC_CSV
    )

    for metric_name, values in etc_values.items():

        summary_rows.append(
            make_summary_row(
                metric_name,
                "",
                values,
            )
        )

    return summary_rows


# =============================================================================
# Excel Styling
# =============================================================================

def apply_outer_border(
    ws,
    min_row,
    max_row,
    min_col,
    max_col,
    border_side,
):
    """Apply one continuous outer border around a rectangular range."""

    for row_index in range(
        min_row,
        max_row + 1,
    ):

        for column_index in range(
            min_col,
            max_col + 1,
        ):

            cell = ws.cell(
                row=row_index,
                column=column_index,
            )

            left = cell.border.left
            right = cell.border.right
            top = cell.border.top
            bottom = cell.border.bottom

            if column_index == min_col:
                left = border_side

            if column_index == max_col:
                right = border_side

            if row_index == min_row:
                top = border_side

            if row_index == max_row:
                bottom = border_side

            cell.border = Border(
                left=left,
                right=right,
                top=top,
                bottom=bottom,
            )


def style_title(ws):
    """Create the report title row."""

    ws.sheet_view.showGridLines = False

    ws.merge_cells(
        "A1:G1"
    )

    title_cell = ws[
        "A1"
    ]

    title_cell.value = (
        "MONTE CARLO CHARACTERIZATION SUMMARY"
    )

    title_cell.font = Font(
        name=FONT_NAME,
        size=16,
        bold=True,
        color=COLOR_WHITE,
    )

    title_cell.fill = PatternFill(
        fill_type="solid",
        fgColor=COLOR_NAVY,
    )

    title_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    ws.row_dimensions[
        1
    ].height = 30


def style_table(
    ws,
    header_row,
    last_row,
):
    """Apply the Monte Carlo summary table style."""

    thin_side = Side(
        style="thin",
        color=COLOR_BORDER,
    )

    outer_side = Side(
        style="medium",
        color=COLOR_SLATE,
    )

    # -------------------------------------------------------------------------
    # Header
    # -------------------------------------------------------------------------

    for column_index in range(
        1,
        8,
    ):

        cell = ws.cell(
            row=header_row,
            column=column_index,
        )

        cell.font = Font(
            name=FONT_NAME,
            size=10,
            bold=True,
            color=COLOR_WHITE,
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=COLOR_SLATE,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    ws.row_dimensions[
        header_row
    ].height = 22

    # -------------------------------------------------------------------------
    # Body
    # -------------------------------------------------------------------------

    for row_index in range(
        header_row + 1,
        last_row + 1,
    ):

        fill_color = (
            COLOR_ROW_ALT
            if (
                row_index
                -
                header_row
            ) % 2 == 0
            else COLOR_WHITE
        )

        fill = PatternFill(
            fill_type="solid",
            fgColor=fill_color,
        )

        for column_index in range(
            1,
            8,
        ):

            cell = ws.cell(
                row=row_index,
                column=column_index,
            )

            cell.fill = fill

            cell.border = Border(
                bottom=thin_side,
            )

            cell.font = Font(
                name=FONT_NAME,
                size=10,
                color=COLOR_TEXT,
            )

            cell.alignment = Alignment(
                vertical="center",
            )

        metric_cell = ws.cell(
            row=row_index,
            column=1,
        )

        metric_cell.font = Font(
            name=FONT_NAME,
            size=10,
            bold=True,
            color=COLOR_NAVY,
        )

        metric_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

        for column_index in (
            2,
            3,
        ):

            ws.cell(
                row=row_index,
                column=column_index,
            ).alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for column_index in range(
            4,
            8,
        ):

            cell = ws.cell(
                row=row_index,
                column=column_index,
            )

            cell.alignment = Alignment(
                horizontal="right",
                vertical="center",
            )

        unit = ws.cell(
            row=row_index,
            column=2,
        ).value

        number_format = (
            "0.000000"
            if unit == "ps"
            else "0.000000E+00"
        )

        for column_index in range(
            4,
            8,
        ):

            cell = ws.cell(
                row=row_index,
                column=column_index,
            )

            if isinstance(
                cell.value,
                (int, float),
            ):

                cell.number_format = (
                    number_format
                )

    # -------------------------------------------------------------------------
    # Table Border
    # -------------------------------------------------------------------------

    apply_outer_border(
        ws,
        min_row=header_row,
        max_row=last_row,
        min_col=1,
        max_col=7,
        border_side=outer_side,
    )

    # -------------------------------------------------------------------------
    # Column Widths
    # -------------------------------------------------------------------------

    fixed_widths = {
        "A": 24,
        "B": 12,
        "C": 10,
        "D": 16,
        "E": 16,
        "F": 16,
        "G": 16,
    }

    for column_letter, width in fixed_widths.items():

        ws.column_dimensions[
            column_letter
        ].width = width

    # -------------------------------------------------------------------------
    # Freeze / Filter
    # -------------------------------------------------------------------------

    ws.freeze_panes = (
        f"A{header_row + 1}"
    )

    ws.auto_filter.ref = (
        f"A{header_row}:G{last_row}"
    )


# =============================================================================
# Workbook Builder
# =============================================================================

def build_workbook(summary_rows):
    """Create the Monte Carlo characterization workbook."""

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "MONTE_SUMMARY"

    style_title(
        worksheet
    )

    header_row = 2

    for column_index, header in enumerate(
        TABLE_HEADERS,
        start=1,
    ):

        worksheet.cell(
            row=header_row,
            column=column_index,
            value=header,
        )

    first_data_row = (
        header_row
        +
        1
    )

    for row_offset, row_values in enumerate(
        summary_rows
    ):

        row_index = (
            first_data_row
            +
            row_offset
        )

        for column_index, value in enumerate(
            row_values,
            start=1,
        ):

            worksheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

    last_row = (
        first_data_row
        +
        len(summary_rows)
        -
        1
    )

    style_table(
        worksheet,
        header_row,
        last_row,
    )

    return workbook


# =============================================================================
# Main
# =============================================================================

def main():

    required_files = [
        SETUP_CSV,
        HOLD_CSV,
        ETC_CSV,
    ]

    for file_path in required_files:

        if not file_path.exists():

            raise FileNotFoundError(
                f"CSV file not found: {file_path}"
            )

    summary_rows = collect_summary_rows()

    workbook = build_workbook(
        summary_rows
    )

    workbook.save(
        OUTPUT_XLSX
    )

    print("")
    print("============================================")
    print(" Monte Carlo Summary Completed")
    print("============================================")
    print("")
    print("Input CSV:")
    print("  ", SETUP_CSV)
    print("  ", HOLD_CSV)
    print("  ", ETC_CSV)
    print("")
    print("Excel created:")
    print("  ", OUTPUT_XLSX)
    print("")


if __name__ == "__main__":
    main()

